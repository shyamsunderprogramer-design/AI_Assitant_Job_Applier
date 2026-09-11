"""Excel tracker tests — run against a temp workbook, no DB or network."""

from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from db.models import Job
from excel.tracker import COL, SHEET_NAME, STATUS_VALUES, ExcelTracker, job_key


def make_job(external_id="1", title="Software Engineer", description="Build things", **kw):
    job = Job(
        source=kw.get("source", "greenhouse"),
        company=kw.get("company", "Acme"),
        company_slug=kw.get("company_slug", "acme"),
        external_id=external_id,
        title=title,
        location=kw.get("location", "Remote - US"),
        description=description,
        requirements=kw.get("requirements", "- Python"),
        application_url=kw.get("application_url", f"https://example.com/jobs/{external_id}"),
        posted_at=kw.get("posted_at", datetime(2026, 8, 1, tzinfo=timezone.utc)),
        content_hash="hash",
        status=kw.get("status", "Not Applied"),
        ats_match_score=kw.get("ats_match_score"),
    )
    job.found_at = datetime(2026, 8, 30, tzinfo=timezone.utc)
    return job


@pytest.fixture
def tracker(tmp_path):
    return ExcelTracker(tmp_path / "tracker.xlsx")


def read(tracker):
    return load_workbook(tracker.path)[SHEET_NAME]


def test_creates_workbook_with_headers(tracker):
    tracker.export([make_job()])
    ws = read(tracker)
    assert ws.cell(row=1, column=COL["Company"]).value == "Company"
    assert ws.cell(row=1, column=COL["ATS Match Score"]).value == "ATS Match Score"


def test_writes_job_values(tracker):
    tracker.export([make_job(title="Backend Engineer")])
    ws = read(tracker)
    assert ws.cell(row=2, column=COL["Company"]).value == "Acme"
    assert ws.cell(row=2, column=COL["Job Title"]).value == "Backend Engineer"
    assert ws.cell(row=2, column=COL["Location"]).value == "Remote - US"
    assert ws.cell(row=2, column=COL["Posting Date"]).value == "2026-08-01"
    assert ws.cell(row=2, column=COL["Application Status"]).value == "Not Applied"


def test_application_link_is_a_hyperlink(tracker):
    tracker.export([make_job()])
    cell = read(tracker).cell(row=2, column=COL["Application Link"])
    assert cell.hyperlink is not None
    assert cell.hyperlink.target == "https://example.com/jobs/1"


def test_appending_does_not_duplicate_existing_rows(tracker):
    job = make_job()
    tracker.export([job])
    appended, updated = tracker.export([job])
    ws = read(tracker)
    assert ws.max_row == 2  # header + one job
    assert (appended, updated) == (0, 1)


def test_new_jobs_append_below_existing(tracker):
    tracker.export([make_job(external_id="1")])
    appended, _ = tracker.export([make_job(external_id="2", title="Data Engineer")])
    ws = read(tracker)
    assert appended == 1
    assert ws.max_row == 3
    assert ws.cell(row=3, column=COL["Job Title"]).value == "Data Engineer"


def test_user_edited_status_survives_reexport(tracker):
    job = make_job()
    tracker.export([job])

    # Simulate the user marking it Applied in Excel.
    wb = load_workbook(tracker.path)
    wb[SHEET_NAME].cell(row=2, column=COL["Application Status"]).value = "Applied"
    wb.save(tracker.path)

    tracker.export([job])
    assert read(tracker).cell(row=2, column=COL["Application Status"]).value == "Applied"


def test_changed_posting_updates_in_place(tracker):
    tracker.export([make_job(title="Engineer I")])
    tracker.export([make_job(title="Engineer II")])
    ws = read(tracker)
    assert ws.max_row == 2
    assert ws.cell(row=2, column=COL["Job Title"]).value == "Engineer II"


def test_score_is_written_when_present(tracker):
    tracker.export([make_job(ats_match_score=0.82)])
    assert read(tracker).cell(row=2, column=COL["ATS Match Score"]).value == 0.82


def test_long_jd_is_truncated_below_excel_cell_limit(tracker):
    tracker.export([make_job(description="x" * 40000)])
    value = read(tracker).cell(row=2, column=COL["JD"]).value
    assert len(value) < 32767
    assert value.endswith("[truncated]")


def test_formatting_applied(tracker):
    tracker.export([make_job()])
    ws = read(tracker)
    assert ws.freeze_panes == "A2"
    assert ws.cell(row=1, column=1).font.bold
    assert ws.column_dimensions["K"].hidden  # Job Key column
    assert ws.auto_filter.ref is not None


def test_status_dropdown_and_conditional_formatting_present(tracker):
    tracker.export([make_job()])
    ws = read(tracker)
    validations = ws.data_validations.dataValidation
    assert validations
    assert "Applied" in validations[0].formula1
    assert len(ws.conditional_formatting._cf_rules) > 0


def test_conditional_formatting_not_duplicated_across_exports(tracker):
    job = make_job()
    tracker.export([job])
    first = len(read(tracker).conditional_formatting._cf_rules)
    tracker.export([make_job(external_id="2")])
    assert len(read(tracker).conditional_formatting._cf_rules) == first


def test_job_key_matches_db_identity():
    assert job_key(make_job(external_id="42")) == "greenhouse:acme:42"


def test_status_values_include_manual_review():
    assert "Manual Review" in STATUS_VALUES


def test_system_flag_reaches_untouched_row(tracker):
    """A row the user hasn't touched accepts a system status change."""
    tracker.export([make_job(status="Not Applied")])
    tracker.export([make_job(status="Manual Review")])
    assert read(tracker).cell(row=2, column=COL["Application Status"]).value == "Manual Review"


def test_system_flag_never_overwrites_user_edit(tracker):
    """Once the user sets a status, no export may change it."""
    tracker.export([make_job(status="Not Applied")])

    wb = load_workbook(tracker.path)
    wb[SHEET_NAME].cell(row=2, column=COL["Application Status"]).value = "Applied"
    wb.save(tracker.path)

    tracker.export([make_job(status="Manual Review")])
    assert read(tracker).cell(row=2, column=COL["Application Status"]).value == "Applied"


def test_stale_manual_review_flag_clears(tracker):
    """Manual Review is system-set, so a later pass may clear it."""
    tracker.export([make_job(status="Manual Review")])
    tracker.export([make_job(status="Not Applied")])
    assert read(tracker).cell(row=2, column=COL["Application Status"]).value == "Not Applied"


def test_user_status_beats_system_status_in_both_directions(tracker):
    tracker.export([make_job(status="Manual Review")])
    wb = load_workbook(tracker.path)
    wb[SHEET_NAME].cell(row=2, column=COL["Application Status"]).value = "Interviewing"
    wb.save(tracker.path)

    tracker.export([make_job(status="Not Applied")])
    assert read(tracker).cell(row=2, column=COL["Application Status"]).value == "Interviewing"


# -- Phase 5: closure reaches the sheet without clobbering the user ---------

def test_closed_status_reaches_an_untouched_row(tracker):
    job = make_job()
    tracker.export([job])

    job.status = "Closed"
    tracker.export([job])

    ws = read(tracker)
    assert ws.cell(row=2, column=COL["Application Status"]).value == "Closed"


def test_closed_status_never_overwrites_a_user_edit(tracker):
    """A posting the user applied to can vanish from the board. Their row stays."""
    job = make_job()
    tracker.export([job])

    ws = read(tracker)
    ws.cell(row=2, column=COL["Application Status"]).value = "Interviewing"
    ws.parent.save(tracker.path)

    job.status = "Closed"
    tracker.export([job])

    assert read(tracker).cell(row=2, column=COL["Application Status"]).value == "Interviewing"


def test_reopened_job_clears_the_closed_status(tracker):
    job = make_job(status="Closed")
    tracker.export([job])

    job.status = "Not Applied"
    tracker.export([job])

    assert read(tracker).cell(row=2, column=COL["Application Status"]).value == "Not Applied"


def test_closed_is_a_selectable_status(tracker):
    assert "Closed" in STATUS_VALUES


# -- Phase 6: prune must clean the sheet, not leave dead rows --------------

def test_remove_deletes_a_pruned_row(tracker):
    tracker.export([make_job("1"), make_job("2")])

    assert tracker.remove({job_key(make_job("2"))}) == 1

    ws = read(tracker)
    keys = {ws.cell(row=r, column=COL["Job Key"]).value for r in range(2, ws.max_row + 1)}
    assert keys == {job_key(make_job("1"))}


def test_remove_never_deletes_a_row_the_user_acted_on(tracker):
    """A posting you applied to can be pruned from the DB. The row stays."""
    tracker.export([make_job("1")])
    ws = read(tracker)
    ws.cell(row=2, column=COL["Application Status"]).value = "Applied"
    ws.parent.save(tracker.path)

    assert tracker.remove({job_key(make_job("1"))}) == 0
    assert read(tracker).cell(row=2, column=COL["Application Status"]).value == "Applied"


def test_remove_reclaims_rows_rather_than_blanking_them(tracker):
    """openpyxl's delete_rows leaves the row's dimensions and styling behind,
    so deleting 111 rows left 111 blank-but-formatted ones and max_row never
    shrank. The body is rebuilt instead."""
    jobs = [make_job(str(i)) for i in range(1, 6)]
    tracker.export(jobs)

    tracker.remove({job_key(j) for j in jobs[:4]})

    ws = read(tracker)
    assert ws.max_row == 2  # header + the one survivor
    assert ws.cell(row=2, column=COL["Job Key"]).value == job_key(jobs[4])


def test_remove_preserves_hyperlinks_on_survivors(tracker):
    tracker.export([make_job("1"), make_job("2")])
    tracker.remove({job_key(make_job("1"))})
    assert read(tracker).cell(row=2, column=COL["Application Link"]).hyperlink is not None


def test_remove_of_an_unknown_key_changes_nothing(tracker):
    tracker.export([make_job("1")])
    assert tracker.remove({"greenhouse:other:99"}) == 0
    assert read(tracker).max_row == 2


def test_remove_with_no_keys_is_a_no_op(tracker):
    tracker.export([make_job("1")])
    assert tracker.remove(set()) == 0
