"""Import company names from a large spreadsheet into a discovery list.

A million-row company export is not a discovery list. Probing all of it would
take decades at polite rates (README.md §C5), and most of it is hopeless
anyway: Greenhouse, Lever and Ashby skew heavily to US tech companies of a
certain size, and the rest of the sheet is banks, hospitals and construction
firms running Workday or Taleo (§C3).

So this filters before it probes, on the three signals that actually predict a
hit:

  * COUNTRY   — match the person's own search locations, not the whole world
  * INDUSTRY  — these ATSs are tech-concentrated
  * HEADCOUNT — below ~50 a company usually has no ATS at all; above ~5,000 it
                is probably on Workday. The band between is where the hits are.

It also prefers the DOMAIN over the company name when one is present. A slug
is far more often the domain root than the trading name — "custom computer
specialists" is at customtech.com, and no amount of name-mangling finds that.
"""

from __future__ import annotations

import csv
import logging
import re
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)

# Industries where Greenhouse/Lever/Ashby actually concentrate.
TECH_INDUSTRIES = (
    "computer software",
    "information technology and service",
    "internet",
    "computer & network security",
    "computer hardware",
    "computer networking",
    "semiconductors",
    "telecommunications",
    "biotechnology",
    "financial services",
    "e-learning",
    "marketing and advertising",
)

# Headcount band worth probing. Outside it the hit rate collapses in both
# directions, for opposite reasons.
DEFAULT_MIN_EMPLOYEES = 50
DEFAULT_MAX_EMPLOYEES = 5000

_DOMAIN_STRIP = re.compile(r"^(?:https?://)?(?:www\.)?", re.I)
_TLD = re.compile(r"\.[a-z.]{2,12}$", re.I)


@dataclass
class CompanyRecord:
    """One row worth probing."""

    name: str
    domain: str | None = None
    industry: str | None = None
    employees: int = 0

    def slug_hint(self) -> str | None:
        """The domain root, which is the single best slug guess available."""
        return domain_root(self.domain)


def domain_root(domain: str | None) -> str | None:
    """'https://www.ClickUp.com/careers' -> 'clickup'."""
    if not domain:
        return None
    cleaned = _DOMAIN_STRIP.sub("", str(domain).strip().lower())
    cleaned = cleaned.split("/")[0].split("?")[0]
    cleaned = _TLD.sub("", cleaned)
    cleaned = cleaned.split(".")[-1] if "." in cleaned else cleaned
    cleaned = re.sub(r"[^a-z0-9-]", "", cleaned)
    if not cleaned or cleaned in JUNK_DOMAINS or len(cleaned) < 2:
        return None
    return cleaned


# Placeholder values that appear in scraped company exports. Probing these
# wastes requests on boards that cannot exist, and a name like "various" would
# otherwise generate a plausible-looking slug.
JUNK_NAMES = {
    "personal", "various", "self employed", "self-employed", "freelance",
    "n/a", "na", "none", "unknown", "confidential", "private", "test",
    "company", "unemployed", "student", "retired", "home", "other",
    "not applicable", "tbd", "stealth", "stealth startup",
}
JUNK_DOMAINS = {
    "websitehere", "example", "domain", "yourcompany", "company", "none",
    "na", "n/a", "test", "website",
}


def is_junk(name: str) -> bool:
    cleaned = re.sub(r"[^a-z0-9 ]", "", (name or "").lower()).strip()
    if len(cleaned) < 2:
        return True
    return cleaned in JUNK_NAMES


def _as_int(value) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def read_spreadsheet(
    path: Path | str,
    sheet: str | None = None,
    countries: tuple[str, ...] = ("united states",),
    industries: tuple[str, ...] = TECH_INDUSTRIES,
    min_employees: int = DEFAULT_MIN_EMPLOYEES,
    max_employees: int = DEFAULT_MAX_EMPLOYEES,
    limit: int = 0,
) -> list[CompanyRecord]:
    """Read and filter a company export (.xlsx or .csv).

    Columns are matched by HEADER NAME, not position, so a differently-ordered
    export still works. Anything it cannot find is simply not filtered on.
    """
    path = Path(path)
    rows = _iter_rows(path, sheet)
    header = next(rows, None)
    if not header:
        return []

    index = _column_index(header)
    if index.get("name") is None:
        raise ValueError(
            f"No company-name column found in {path.name}. "
            f"Looked for one of: name, company, company name. Got: {header}"
        )

    wanted_countries = {c.lower() for c in countries} if countries else set()
    records: list[CompanyRecord] = []
    seen: set[str] = set()

    for row in rows:
        name = _cell(row, index.get("name"))
        if not name or is_junk(name):
            continue

        if wanted_countries:
            country = (_cell(row, index.get("country")) or "").lower()
            if country and country not in wanted_countries:
                continue

        industry = (_cell(row, index.get("industry")) or "").lower()
        if industries and industry and not any(t in industry for t in industries):
            continue

        employees = _as_int(_cell(row, index.get("employees")))
        # Only filter on headcount when the sheet actually states one —
        # a missing value is unknown, not zero.
        if employees:
            if min_employees and employees < min_employees:
                continue
            if max_employees and employees > max_employees:
                continue

        key = name.strip().lower()
        if key in seen:
            continue
        seen.add(key)

        records.append(
            CompanyRecord(
                name=name.strip(),
                domain=_cell(row, index.get("domain")),
                industry=industry or None,
                employees=employees,
            )
        )
        if limit and len(records) >= limit:
            break

    # Bigger companies are likelier to run a real ATS, so probe them first.
    # A run that is cut short should have spent its requests on the best bets.
    records.sort(key=lambda r: r.employees, reverse=True)
    return records


def _iter_rows(path: Path, sheet: str | None):
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        for row in ws.iter_rows(values_only=True):
            yield list(row)
        return

    with open(path, newline="", encoding="utf-8", errors="replace") as fh:
        for row in csv.reader(fh):
            yield row


# Header aliases, lowercased and stripped of punctuation.
_ALIASES = {
    "name": ("company name", "company", "name", "organization", "employer"),
    "industry": ("industry", "sector", "category"),
    "country": ("country", "hq country"),
    "domain": ("website domain", "website", "domain", "url", "company domain"),
    "employees": ("employees", "employee count", "headcount", "size", "company size"),
}


def _column_index(header: list) -> dict[str, int | None]:
    normalised = [re.sub(r"[^a-z0-9 ]", " ", str(h or "").lower()).strip() for h in header]
    normalised = [re.sub(r"\s+", " ", h) for h in normalised]

    found: dict[str, int | None] = {}
    for field, aliases in _ALIASES.items():
        found[field] = next((i for i, h in enumerate(normalised) if h in aliases), None)
        if found[field] is None:
            found[field] = next(
                (i for i, h in enumerate(normalised) if any(a in h for a in aliases)), None
            )
    return found


def _cell(row: list, index: int | None) -> str | None:
    if index is None or index >= len(row):
        return None
    value = row[index]
    return None if value is None else str(value).strip() or None


def write_names_file(records: list[CompanyRecord], path: Path | str) -> Path:
    """Write a discovery input file: `name` or `name,domain` per line."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Generated from a company export — see scraper/company_import.py",
        "# Format: company name[,domain]. The domain is the better slug guess.",
        "",
    ]
    for record in records:
        lines.append(f"{record.name},{record.domain}" if record.domain else record.name)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path
