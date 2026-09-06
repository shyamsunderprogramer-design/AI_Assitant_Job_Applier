"""Base-resume parsing. Reads .docx or .pdf into a plain-text structure.

Deliberately simple: the goal is faithful text extraction, not layout fidelity.
Everything downstream (tailoring, scoring, the fabrication guard) works off the
text, and the guard depends on this text being *complete* — if a real skill is
missing here it will look like an invention later.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger(__name__)

BULLET_CHARS = ("-", "•", "▪", "◦", "*", "‣", "–", "—")

# Headings we expect in a resume; used to segment the document.
KNOWN_HEADINGS = (
    "summary", "objective", "profile", "about",
    "experience", "work experience", "professional experience", "employment",
    "education", "skills", "technical skills", "core competencies",
    "projects", "certifications", "publications", "awards", "achievements",
    "volunteer", "interests", "languages", "references",
)


@dataclass
class Section:
    heading: str
    lines: list[str] = field(default_factory=list)

    @property
    def bullets(self) -> list[str]:
        return [ln for ln in self.lines if is_bullet(ln)]

    def text(self) -> str:
        return "\n".join([self.heading, *self.lines]).strip()


@dataclass
class Resume:
    raw_text: str
    sections: list[Section] = field(default_factory=list)
    source_path: Path | None = None

    def section(self, name: str) -> Section | None:
        target = name.lower()
        return next((s for s in self.sections if target in s.heading.lower()), None)

    def all_bullets(self) -> list[str]:
        return [b for s in self.sections for b in s.bullets]

    def text(self) -> str:
        return self.raw_text


def is_bullet(line: str) -> bool:
    stripped = line.strip()
    return bool(stripped) and stripped.startswith(BULLET_CHARS)


def strip_bullet(line: str) -> str:
    stripped = line.strip()
    for char in BULLET_CHARS:
        if stripped.startswith(char):
            return stripped[len(char) :].strip()
    return stripped


def looks_like_heading(line: str) -> bool:
    """A short line that is all-caps or matches a known resume section name."""
    stripped = line.strip().rstrip(":")
    if not stripped or len(stripped) > 48 or is_bullet(line):
        return False
    if stripped.lower() in KNOWN_HEADINGS:
        return True
    # ALL CAPS with no sentence punctuation, e.g. "PROFESSIONAL EXPERIENCE"
    letters = [c for c in stripped if c.isalpha()]
    return bool(letters) and all(c.isupper() for c in letters) and "." not in stripped


def parse_docx(path: Path | str) -> Resume:
    from docx import Document  # imported lazily so Phase 1/2 don't need python-docx

    document = Document(str(path))
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        # Word list styles carry the bullet as formatting, not as a character.
        if para.style is not None and "list" in (para.style.name or "").lower():
            text = text if is_bullet(text) else f"- {text}"
        lines.append(text)

    # Tables are ATS-hostile but people still use them; pull their text in too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    lines.extend(ln.strip() for ln in cell_text.split("\n") if ln.strip())

    return _build(lines, Path(path))


def parse_pdf(path: Path | str) -> Resume:
    from pypdf import PdfReader  # lazy import

    reader = PdfReader(str(path))
    lines: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        for line in text.split("\n"):
            stripped = line.strip()
            if stripped:
                lines.append(stripped)
    return _build(lines, Path(path))


def parse_resume(path: Path | str) -> Resume:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Base resume not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix in (".txt", ".md"):
        return _build([ln.strip() for ln in path.read_text(encoding="utf-8").split("\n") if ln.strip()], path)
    raise ValueError(f"Unsupported resume format: {suffix} (use .docx or .pdf)")


def find_base_resume(directory: Path | str) -> Path | None:
    """Locate the user's base resume in resume/, whatever they named it."""
    directory = Path(directory)
    if not directory.exists():
        return None
    candidates = [
        p
        for p in directory.iterdir()
        if p.is_file()
        and p.suffix.lower() in (".docx", ".pdf", ".txt", ".md")
        and not p.name.startswith("~$")  # Word lock files
    ]
    if not candidates:
        return None
    # Prefer a file that says "base", else the most recently modified.
    preferred = [p for p in candidates if "base" in p.stem.lower()]
    pool = preferred or candidates
    return max(pool, key=lambda p: p.stat().st_mtime)


def _build(lines: list[str], path: Path | None) -> Resume:
    sections: list[Section] = []
    current = Section(heading="HEADER")

    for line in lines:
        if looks_like_heading(line):
            if current.lines or current.heading != "HEADER":
                sections.append(current)
            current = Section(heading=line.strip().rstrip(":"))
        else:
            current.lines.append(line)

    sections.append(current)
    return Resume(
        raw_text="\n".join(lines),
        sections=[s for s in sections if s.lines or s.heading != "HEADER"],
        source_path=path,
    )
