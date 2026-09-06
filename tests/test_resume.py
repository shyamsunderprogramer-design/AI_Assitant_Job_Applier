"""Phase 3 tests — all offline. No API key, no network, no cost."""

import pytest
from docx import Document

from resume.guard import check_no_fabrication
from resume.parser import find_base_resume, is_bullet, looks_like_heading, parse_resume, strip_bullet
from resume.scorer import content_terms, jd_keyword_weights, score_resume
from resume.tailor import TailorResult, _extract_json
from resume.writer import output_filename, write_review_note, write_tailored_resume

BASE_RESUME = """JANE DOE
jane@example.com | 555-0100 | San Francisco, CA

SUMMARY
Backend engineer with 6 years building Python services.

EXPERIENCE
Acme Corp — Senior Software Engineer, 2021-2026
- Built REST APIs in Python and Django serving 40000 requests per day
- Migrated the billing service to PostgreSQL, cutting query time by 35%
- Mentored 3 junior engineers on code review practices

Globex — Software Engineer, 2019-2021
- Developed data pipelines with Airflow and SQL

SKILLS
Python, Django, PostgreSQL, SQL, Airflow, Git, Docker

EDUCATION
BS Computer Science, State University, 2019
"""

JD = """We are hiring a Senior Backend Engineer.
You will build scalable Python services and REST APIs.
Requirements:
- Strong Python and Django experience
- PostgreSQL and SQL expertise
- Experience with Kubernetes and distributed systems
- Airflow or similar orchestration tooling
"""


# --- scorer ---------------------------------------------------------------

def test_content_terms_drops_stopwords_and_digits():
    terms = content_terms("We are hiring a Python engineer with 5 years")
    assert "python" in terms
    assert "the" not in terms and "are" not in terms
    assert "5" not in terms


def test_jd_weights_rank_repeated_terms_higher():
    weights = jd_keyword_weights("python python python kubernetes")
    assert weights["python"] > weights["kubernetes"]


def test_score_rewards_overlap():
    strong = score_resume(BASE_RESUME, JD)
    weak = score_resume("Graphic designer skilled in Photoshop and Illustrator", JD)
    assert strong.score > weak.score
    assert 0.0 <= strong.score <= 1.0


def test_score_reports_real_gaps():
    result = score_resume(BASE_RESUME, JD)
    assert "kubernetes" in result.missing      # genuinely absent from the resume
    assert "python" in result.matched
    assert "django" in result.matched


def test_score_handles_empty_jd():
    result = score_resume(BASE_RESUME, "")
    assert result.score == 0.0
    assert result.jd_terms == 0


def test_percent_and_summary_render():
    result = score_resume(BASE_RESUME, JD)
    assert 0 <= result.percent <= 100
    assert "match" in result.summary()


# --- fabrication guard ----------------------------------------------------

def test_guard_accepts_faithful_rephrasing():
    tailored = "Built REST APIs using Python and Django handling 40000 requests daily"
    assert check_no_fabrication(BASE_RESUME, tailored).ok


def test_guard_catches_invented_skill():
    result = check_no_fabrication(BASE_RESUME, "Deployed services on Kubernetes at scale")
    assert not result.ok
    assert any(v.kind == "skill" and v.value == "kubernetes" for v in result.violations)


def test_guard_catches_invented_metric():
    result = check_no_fabrication(BASE_RESUME, "Cut latency by 92% across the platform")
    assert not result.ok
    assert any(v.kind == "metric" for v in result.violations)


def test_guard_catches_invented_employer():
    result = check_no_fabrication(BASE_RESUME, "Led platform work at Initech")
    assert not result.ok
    assert any(v.kind == "entity" and v.value == "Initech" for v in result.violations)


def test_guard_catches_invented_year():
    result = check_no_fabrication(BASE_RESUME, "Worked on the billing service in 2014")
    assert not result.ok
    assert any(v.kind == "year" and v.value == "2014" for v in result.violations)


def test_guard_allows_real_metrics_and_employers():
    tailored = "At Acme Corp, reduced query time 35% after the PostgreSQL migration"
    assert check_no_fabrication(BASE_RESUME, tailored).ok


def test_guard_ignores_common_action_verbs():
    assert check_no_fabrication(BASE_RESUME, "Delivered Python services with Django").ok


def test_guard_deduplicates_repeat_violations():
    tailored = "Used Kubernetes here.\nAlso used Kubernetes there."
    violations = check_no_fabrication(BASE_RESUME, tailored).violations
    assert len([v for v in violations if v.value == "kubernetes"]) == 1


def test_guard_report_is_readable():
    report = check_no_fabrication(BASE_RESUME, "Ran Kubernetes clusters").report()
    assert "kubernetes" in report.lower()


# --- parser ---------------------------------------------------------------

def test_bullet_helpers():
    assert is_bullet("- did a thing")
    assert is_bullet("• did a thing")
    assert not is_bullet("Did a thing")
    assert strip_bullet("- did a thing") == "did a thing"


def test_heading_detection():
    assert looks_like_heading("EXPERIENCE")
    assert looks_like_heading("Skills")
    assert not looks_like_heading("- Built REST APIs in Python")
    assert not looks_like_heading("A rather long line that is clearly body text, not a heading")


def test_parse_txt_resume_sections(tmp_path):
    path = tmp_path / "base_resume.txt"
    path.write_text(BASE_RESUME, encoding="utf-8")
    resume = parse_resume(path)
    assert resume.section("EXPERIENCE") is not None
    assert resume.section("SKILLS") is not None
    assert len(resume.all_bullets()) == 4


def test_parse_docx_roundtrip(tmp_path):
    doc = Document()
    doc.add_paragraph("JANE DOE")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("- Built REST APIs in Python")
    path = tmp_path / "base.docx"
    doc.save(path)

    resume = parse_resume(path)
    assert "Built REST APIs in Python" in resume.text()
    assert resume.section("EXPERIENCE") is not None


def test_unsupported_format_rejected(tmp_path):
    path = tmp_path / "resume.rtf"
    path.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        parse_resume(path)


def test_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        parse_resume(tmp_path / "nope.docx")


def test_find_base_resume_prefers_base_named_file(tmp_path):
    (tmp_path / "old_cv.pdf").write_bytes(b"x")
    (tmp_path / "base_resume.docx").write_bytes(b"x")
    assert find_base_resume(tmp_path).name == "base_resume.docx"


def test_find_base_resume_ignores_word_lock_files(tmp_path):
    (tmp_path / "~$base_resume.docx").write_bytes(b"x")
    assert find_base_resume(tmp_path) is None


# --- tailor helpers -------------------------------------------------------

def test_extract_json_handles_code_fence():
    payload = _extract_json('```json\n{"summary": "hi", "bullets": []}\n```')
    assert payload["summary"] == "hi"


def test_extract_json_handles_bare_object():
    assert _extract_json('{"gaps": ["kubernetes"]}')["gaps"] == ["kubernetes"]


def test_extract_json_rejects_prose():
    with pytest.raises(ValueError):
        _extract_json("I could not complete that request.")


def test_tailor_result_accepted_reflects_guard():
    result = TailorResult(summary="Built Python services", bullets=[])
    result.guard = check_no_fabrication(BASE_RESUME, result.tailored_text())
    assert result.accepted

    bad = TailorResult(summary="Ran Kubernetes clusters", bullets=[])
    bad.guard = check_no_fabrication(BASE_RESUME, bad.tailored_text())
    assert not bad.accepted


# --- writer ---------------------------------------------------------------

def test_output_filename_is_deterministic_and_safe():
    name = output_filename("Acme Corp!", "Senior Software Engineer, Platform", "4567")
    assert name == "acme-corp_senior-software-engineer-platform_4567.docx"
    assert name == output_filename("Acme Corp!", "Senior Software Engineer, Platform", "4567")


def test_written_docx_is_ats_safe(tmp_path):
    resume = parse_resume(_write_base(tmp_path))
    result = TailorResult(
        summary="Backend engineer with 6 years building Python services.",
        bullets=[{
            "original": "- Built REST APIs in Python and Django serving 40000 requests per day",
            "tailored": "Built scalable REST APIs in Python and Django",
            "reason": "matches JD language",
        }],
    )
    out = write_tailored_resume(resume, result, tmp_path / "out.docx")

    doc = Document(str(out))
    assert len(doc.tables) == 0                      # no tables
    assert len(doc.inline_shapes) == 0               # no images
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Built scalable REST APIs" in text        # rewrite applied
    assert "Acme Corp" in text                       # employer preserved verbatim
    assert "State University" in text                # education preserved


def test_writer_preserves_untailored_bullets(tmp_path):
    resume = parse_resume(_write_base(tmp_path))
    out = write_tailored_resume(resume, TailorResult(summary=None), tmp_path / "out2.docx")
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "Mentored 3 junior engineers" in text


def test_review_note_lists_gaps_and_rejections(tmp_path):
    result = TailorResult(summary=None, gaps=["kubernetes"])
    result.guard = check_no_fabrication(BASE_RESUME, "Ran Kubernetes clusters")
    path = write_review_note(result, tmp_path / "note.txt")
    body = path.read_text(encoding="utf-8")
    assert "kubernetes" in body.lower()
    assert "REJECTED" in body


def _write_base(tmp_path):
    path = tmp_path / "base_resume.txt"
    path.write_text(BASE_RESUME, encoding="utf-8")
    return path
